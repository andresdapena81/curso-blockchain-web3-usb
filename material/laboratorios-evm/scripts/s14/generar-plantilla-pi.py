"""
Laboratorio 14 · genera la hoja de cálculo de pérdida impermanente

    python scripts/s14/generar-plantilla-pi.py                 # plantilla + CSV en guias/
    python scripts/s14/generar-plantilla-pi.py --solucion RUTA  # además, la versión resuelta (docente)

Requiere openpyxl (python -m pip install openpyxl). El estudiante NO necesita
correr esto: recibe ya generada guias/s14-plantilla-perdida-impermanente.xlsx.

La plantilla trae las celdas amarillas VACÍAS (ahí van sus fórmulas) y, debajo,
los VALORES ESPERADOS calculados aquí mismo, con una columna que dice si lo
suyo coincide. Los valores esperados salen de la misma aritmética que
scripts/s14/perdida-impermanente.js; si cambian allá, cambian aquí.
"""

import csv
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAIZ = Path(__file__).resolve().parents[2]
DESTINO = RAIZ / "guias" / "s14-plantilla-perdida-impermanente.xlsx"
DESTINO_CSV = RAIZ / "guias" / "s14-plantilla-perdida-impermanente.csv"

X0, Y0, P0 = 1, 2000, 2000
ESCENARIOS = [("A · ETH cae a 1 000", 1000), ("B · ETH sube a 4 000", 4000), ("C · ETH sube a 8 000", 8000)]

TINTA = "0F0F11"
NARANJA = "F07F06"
AMARILLO = PatternFill("solid", fgColor="FFF4C2")
GRIS = PatternFill("solid", fgColor="F2F0F5")
NEGRO = PatternFill("solid", fgColor=TINTA)
BORDE = Border(*(Side(style="thin", color="D9D6E0"),) * 4)
BLANCO_B = Font(name="Arial", bold=True, color="FFFFFF")
B = Font(name="Arial", bold=True)
N = Font(name="Arial")


def esperado(p1):
    r = p1 / P0
    eth = X0 / math.sqrt(r)
    usd = Y0 * math.sqrt(r)
    pool = eth * p1 + usd
    cons = X0 * p1 + Y0
    return r, eth, usd, pool, cons, 1 - pool / cons


def formulas(fila):
    """Las fórmulas de la solución para la fila `fila` de la tabla de trabajo."""
    return {
        "C": f"=B{fila}/$B$5",
        "D": f"=$B$3/SQRT(C{fila})",
        "E": f"=$B$4*SQRT(C{fila})",
        "F": f"=D{fila}*B{fila}+E{fila}",
        "G": f"=$B$3*B{fila}+$B$4",
        "H": f"=1-F{fila}/G{fila}",
    }


def construir(resuelta: bool) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja de trabajo"
    ws.sheet_view.showGridLines = False
    for col, ancho in zip("ABCDEFGHI", [26, 14, 10, 14, 14, 14, 16, 12, 16]):
        ws.column_dimensions[col].width = ancho

    ws["A1"] = "LABORATORIO 14 · PÉRDIDA IMPERMANENTE EN TRES ESCENARIOS"
    ws["A1"].font = Font(name="Arial Black", size=13, color=TINTA)
    ws["A2"] = "Llenen las celdas AMARILLAS con FÓRMULAS (no con números copiados). La columna I dice si coinciden con lo esperado."
    ws["A2"].font = Font(name="Arial", italic=True, color="6F6B80")

    datos = [("ETH depositado (x0)", X0), ("USDC depositado (y0)", Y0), ("Precio inicial del ETH (P0)", P0)]
    for i, (et, v) in enumerate(datos, start=3):
        ws[f"A{i}"], ws[f"B{i}"] = et, v
        ws[f"A{i}"].font, ws[f"B{i}"].font, ws[f"B{i}"].fill = B, N, GRIS
    ws["A6"], ws["B6"] = "k = x0 · y0", "=B3*B4"
    ws["A6"].font = B

    cab = ["Escenario", "Precio final P1", "r = P1/P0", "ETH en el pool", "USDC en el pool",
           "Valor del pool", "Valor si conserva", "Pérdida", "¿Coincide?"]
    fila_cab = 8
    for j, t in enumerate(cab):
        c = ws.cell(row=fila_cab, column=j + 1, value=t)
        c.font, c.fill, c.alignment = BLANCO_B, NEGRO, Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[fila_cab].height = 30

    fila_esp0 = 16   # primera fila de la tabla de valores esperados
    filas = ESCENARIOS + [("D · su propio escenario", None)]
    for i, (nombre, p1) in enumerate(filas):
        f = fila_cab + 1 + i
        ws.cell(row=f, column=1, value=nombre).font = B
        propio = p1 if p1 else (P0 * 10 if resuelta else None)   # la solución usa r = 10
        ws.cell(row=f, column=2, value=propio).fill = GRIS if p1 else AMARILLO
        fx = formulas(f)
        for col in "CDEFGH":
            celda = ws[f"{col}{f}"]
            celda.fill = AMARILLO
            if resuelta:
                celda.value = fx[col]
        for col in "BCDEFGH":
            ws[f"{col}{f}"].border = BORDE
        ws[f"C{f}"].number_format = "0.00"
        ws[f"D{f}"].number_format = "0.0000"
        for col in "EFG":
            ws[f"{col}{f}"].number_format = "#,##0.00"
        ws[f"H{f}"].number_format = "0.00%"
        if p1:
            e = fila_esp0 + i
            ws[f"I{f}"] = (f'=IF(AND(ABS(C{f}-C{e})<0.0001,ABS(D{f}-D{e})<0.0001,ABS(E{f}-E{e})<0.01,'
                           f'ABS(F{f}-F{e})<0.01,ABS(G{f}-G{e})<0.01,ABS(H{f}-H{e})<0.0001),"✔ coincide","✘ revisar")')
        else:
            # el escenario propio se contrasta con la fórmula cerrada 1 − 2√r/(1+r)
            ws[f"I{f}"] = f'=IF(ISNUMBER(C{f}),IF(ABS(H{f}-(1-2*SQRT(C{f})/(1+C{f})))<0.0001,"✔ coincide","✘ revisar"),"")'
        ws[f"I{f}"].font = B

    ws.cell(row=fila_esp0 - 2, column=1, value="VALORES ESPERADOS (para autocomprobarse; no los copien arriba: arriba van fórmulas)").font = Font(name="Arial", bold=True, color=NARANJA)
    for j, t in enumerate(cab[:-1]):
        c = ws.cell(row=fila_esp0 - 1, column=j + 1, value=t)
        c.font, c.fill = BLANCO_B, NEGRO
    for i, (nombre, p1) in enumerate(ESCENARIOS):
        f = fila_esp0 + i
        r, eth, usd, pool, cons, pi = esperado(p1)
        for col, v, fmt in [("A", nombre, None), ("B", p1, "#,##0"), ("C", round(r, 6), "0.00"),
                            ("D", round(eth, 6), "0.0000"), ("E", round(usd, 4), "#,##0.00"),
                            ("F", round(pool, 4), "#,##0.00"), ("G", round(cons, 4), "#,##0.00"),
                            ("H", round(pi, 6), "0.00%")]:
            ws[f"{col}{f}"] = v
            ws[f"{col}{f}"].border = BORDE
            if fmt:
                ws[f"{col}{f}"].number_format = fmt

    notas = [
        "Pistas: r = P1/P0 · ETH en el pool = x0/√r · USDC en el pool = y0·√r · valor del pool = ETH·P1 + USDC",
        "valor si conserva = x0·P1 + y0 · pérdida = 1 − valor del pool / valor si conserva.",
        "Comprobación cruzada: la pérdida debe ser igual a 1 − 2·√r/(1+r) (lámina A.7 y scripts/s14/perdida-impermanente.js).",
        "Preguntas para el informe: ¿por qué A y B dan el MISMO porcentaje? ¿qué comisión anual tendría que ganar el LP en C para no perder?",
    ]
    for i, t in enumerate(notas):
        ws.cell(row=21 + i, column=1, value=t).font = Font(name="Arial", size=9, color="3A3844")
    return wb


def main():
    construir(False).save(DESTINO)
    with DESTINO_CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["escenario", "precio_final", "r", "eth_en_pool", "usdc_en_pool",
                    "valor_pool", "valor_conservar", "perdida_impermanente"])
        for nombre, p1 in ESCENARIOS:
            r, eth, usd, pool, cons, pi = esperado(p1)
            w.writerow([nombre, p1, f"{r:.4f}", f"{eth:.6f}", f"{usd:.4f}", f"{pool:.4f}", f"{cons:.4f}", f"{pi:.6f}"])
    print(f"plantilla  → {DESTINO}")
    print(f"csv        → {DESTINO_CSV}")
    if "--solucion" in sys.argv:
        ruta = Path(sys.argv[sys.argv.index("--solucion") + 1])
        ruta.parent.mkdir(parents=True, exist_ok=True)
        construir(True).save(ruta)
        print(f"solución   → {ruta}")


if __name__ == "__main__":
    main()
